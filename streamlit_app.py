# modern_grades_dashboard.py

import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
from statistics import mode, StatisticsError
import re
import tempfile
import os

from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from PIL import Image as PILImage
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# --- Nauja: interaktyvioms diagramoms
import plotly.express as px
import plotly.graph_objects as go

st.set_page_config(page_title="Mokinių pasiekimų analizė", layout="wide")

# --- Minimalus šiuolaikinis stilius
st.markdown("""
<style>
    .block-container { padding-top: 1rem; padding-bottom: 2rem; }
    h1, h2, h3, h4 { font-family: 'Segoe UI', system-ui, -apple-system, sans-serif; color: #2F4F4F; }
    .metric-row .stMetric { border-radius: 16px; padding: 0.5rem 1rem; background: #f8fafc; box-shadow: 0 1px 3px rgba(0,0,0,0.06); }
    .stTabs [data-baseweb="tab-list"] { gap: .25rem; }
    .stTabs [data-baseweb="tab"] { background: #f1f5f9; border-radius: 12px; padding: .5rem 1rem; }
    .stTabs [aria-selected="true"] { background: #e2e8f0; }
</style>
""", unsafe_allow_html=True)

st.title("📘 Mokinių pasiekimų analizė")

uploaded_file = st.file_uploader("📂 Įkelkite Excel (.xlsx) failą su suvestine", type=["xlsx"])

if uploaded_file:
    # 🔁 Išsisaugome originalų failą baitais (kad vėliau galėtume jį atkurti ir nuspalvinti)
    uploaded_bytes = uploaded_file.read()
    excel_buffer_for_pd = BytesIO(uploaded_bytes)

    # --- NUSKAITYMAS
    df_raw = pd.read_excel(excel_buffer_for_pd, sheet_name="Pasiekimų ir lankomumo", skiprows=3)
    df_raw = df_raw.rename(columns={df_raw.columns[0]: "Eil_Nr", df_raw.columns[1]: "Vardas_Pavarde"})
    df_raw = df_raw[df_raw["Vardas_Pavarde"].notna() & df_raw["Eil_Nr"].notna()]

    attendance_cols = [
        "Praleistos pamokos",
        "Pateisintos dėl ligos",
        "Pateisintos dėl kitų priežasčių",
        "Nepateisinta"
    ]
    attendance_df = df_raw[["Vardas_Pavarde"] + [col for col in attendance_cols if col in df_raw.columns]].copy()

    # Stulpeliai su pažymiais
    subject_columns = [col for col in df_raw.columns[2:25]
                       if all(x not in str(col).lower() for x in ["vidurkis", "metinis", "pusm", "socialinė"])]

    def extract_numeric(x):
        match = re.search(r"(\d+)", str(x))
        return int(match.group(1)) if match else None

    df_raw[subject_columns] = df_raw[subject_columns].map(extract_numeric)

    # Long formato rinkinys
    df_long = df_raw.melt(
        id_vars=["Vardas_Pavarde"],
        value_vars=subject_columns,
        var_name="Dalykas",
        value_name="Ivertinimas"
    )
    df_long = df_long[df_long["Ivertinimas"].notna()]

    # --- ŠONINĖ JUOSTA: FILTRAI
    st.sidebar.header("🎛️ Filtravimas")
    # Paieška pagal vardą
    search = st.sidebar.text_input("🔎 Paieška pagal vardą (apytikslė)", "")
    # Daugialypis pasirinkimas
    students_all = sorted(df_long["Vardas_Pavarde"].unique())
    subjects_all = sorted(df_long["Dalykas"].unique())
    selected_students = st.sidebar.multiselect("👥 Mokiniai", ["Visi"] + students_all, default=["Visi"])
    selected_subjects = st.sidebar.multiselect("📚 Dalykai", ["Visi"] + subjects_all, default=["Visi"])
    # Pažymių diapazonas
    grade_min, grade_max = st.sidebar.slider("📈 Pažymių diapazonas", 1, 10, (1, 10))

    # --- Pritaikyti filtrus
    filtered_df = df_long.copy()
    # Paieška (case-insensitive contains)
    if search.strip():
        s = search.strip().casefold()
        filtered_df = filtered_df[filtered_df["Vardas_Pavarde"].str.casefold().str.contains(s, na=False)]

    # Mokiniai
    if selected_students and "Visi" not in selected_students:
        filtered_df = filtered_df[filtered_df["Vardas_Pavarde"].isin(selected_students)]
    # Dalykai
    if selected_subjects and "Visi" not in selected_subjects:
        filtered_df = filtered_df[filtered_df["Dalykas"].isin(selected_subjects)]
    # Pažymių diapazonas
    filtered_df = filtered_df[(filtered_df["Ivertinimas"] >= grade_min) & (filtered_df["Ivertinimas"] <= grade_max)]

    # --- KPI kortelės
    mean_score = filtered_df["Ivertinimas"].mean()
    std_score = filtered_df["Ivertinimas"].std()
    try:
        moda_score = mode(filtered_df["Ivertinimas"])
    except StatisticsError:
        moda_score = "—"

    record_count = len(filtered_df)
    student_count = filtered_df["Vardas_Pavarde"].nunique()

    st.markdown('<div class="metric-row">', unsafe_allow_html=True)
    col1, col2, col3, col4, col5 = st.columns(5)
    col1.metric("📊 Vidurkis", f"{mean_score:.2f}" if pd.notna(mean_score) else "—")
    col2.metric("📌 Moda", moda_score if moda_score != "—" else "—")
    col3.metric("📉 Standartinis nuokrypis", f"{std_score:.2f}" if pd.notna(std_score) else "—")
    col4.metric("🧾 Įrašų skaičius", record_count)
    col5.metric("👥 Mokinių skaičius", student_count)
    st.markdown('</div>', unsafe_allow_html=True)

    # --- FUNKCIJOS: lygio skaičiavimas pagal MIN pažymį
    def student_overall_level(grades):
        """Mokinio lygis pagal MAŽIAUSIĄ jo pažymį pasirinktame kontekste."""
        g = [int(x) for x in grades if pd.notna(x)]
        if not g:
            return "—"
        mn = min(g)
        if mn >= 9:   return "Aukštesnysis (9–10)"
        if mn >= 7:   return "Pagrindinis (7–10)"
        if mn >= 5:   return "Patenkinamas (5–10)"
        if mn >= 4:   return "Slenkstinis (4–10)"
        return "Nepatenkinamas (1–10)"
    # --- Spalvų žemėlapis pasiekimų lygiams (naudojamas grafikuose ir lentelėse)
    level_color_map = {
        "Aukštesnysis (9–10)": "#86efac",   # žalia
        "Pagrindinis (7–10)": "#fde68a",    # geltona
        "Patenkinamas (5–10)": "#bfdbfe",   # mėlyna
        "Slenkstinis (4–10)": "#fca5a5",    # šv. raudona
        "Nepatenkinamas (1–10)": "#f87171", # ryški raudona
        "Žalia": "#c7f1c0",
        "Geltona": "#fff3b0",
        "Raudona": "#f8c8c8"
    }

    # --- SKIRTUKAI
    tab_overview, tab_subjects, tab_students, tab_levels, tab_exports, tab_assistant = st.tabs(
        ["🔎 Apžvalga", "📚 Dalykai", "👤 Mokiniai", "📊 Lygiai", "📤 Eksportas","asistentas"]
    )

    # === 🔎 APŽVALGA
    with tab_overview:
        # Interaktyvus "Vidurkiai pagal dalyką"
        subject_avg = filtered_df.groupby("Dalykas")["Ivertinimas"].mean().sort_values()
        fig_subjects = px.bar(
            subject_avg.reset_index(),
            x="Ivertinimas", y="Dalykas", orientation="h",
            title="Vidutiniai įvertinimai pagal dalyką",
            text="Ivertinimas"
        )
        fig_subjects.update_traces(texttemplate="%{text:.1f}", textposition="outside")
        fig_subjects.update_layout(xaxis_range=[0, 10], margin=dict(l=10, r=20, t=50, b=10), height=500)
        st.plotly_chart(fig_subjects, use_container_width=True)

        # Interaktyvus "Mokiniai pagal bendrą vidurkį"
        student_avg = filtered_df.groupby("Vardas_Pavarde")["Ivertinimas"].mean().sort_values(ascending=True)
        fig_students = px.bar(
            student_avg.reset_index(),
            x="Ivertinimas", y="Vardas_Pavarde", orientation="h",
            title="Mokiniai pagal bendrą vidurkį",
            text="Ivertinimas"
        )
        fig_students.update_traces(texttemplate="%{text:.1f}", textposition="outside")
        fig_students.update_layout(xaxis_range=[0, 10], margin=dict(l=10, r=20, t=50, b=10), height=600)
        st.plotly_chart(fig_students, use_container_width=True)

        # Interaktyvi lentelė
        st.markdown("#### 📋 Įrašai (filtruoti)")
        st.dataframe(filtered_df.sort_values(["Vardas_Pavarde", "Dalykas"]), use_container_width=True)
        st.download_button(
            "⬇️ Atsisiųsti filtruotus duomenis (CSV)",
            data=filtered_df.to_csv(index=False).encode("utf-8"),
            file_name="filtruoti_duomenys.csv",
            mime="text/csv"
        )

    # === 📚 DALYKAI
    with tab_subjects:
        st.subheader("Dalykų vidurkiai ir sklaida")
        # Box plot per dalyką (interaktyvus)
        fig_box = px.box(
            filtered_df, x="Dalykas", y="Ivertinimas", points="all",
            title="Pažymių sklaida pagal dalyką"
        )
        fig_box.update_layout(xaxis_tickangle=-30, margin=dict(l=10, r=20, t=50, b=80), height=550)
        st.plotly_chart(fig_box, use_container_width=True)

        # Pivot tipo lentelė
        pivot = filtered_df.pivot_table(index="Dalykas", values="Ivertinimas", aggfunc=["count", "mean", "min", "max"])
        pivot.columns = ["Kiekis", "Vidurkis", "Min", "Max"]
        pivot = pivot.sort_values("Vidurkis", ascending=False).reset_index()
        st.dataframe(pivot, use_container_width=True)

        # === 👤 MOKINIAI — tik spalvota lentelė pagal pasiektą lygį (Žalia/Geltona/Raudona)
    with tab_students:
        st.subheader("Mokinių santrauka (spalvinis žymėjimas)")

        # Skaičiuojame pagal VISUS mokinio pažymius (nepaisant aktyvių filtrų),
        # kaip ir prašėte anksčiau.
        grade_lists = df_long.groupby("Vardas_Pavarde")["Ivertinimas"].apply(list).reset_index(name="Visi_iverciai")


        def student_flag(grades):
            grades = [g for g in grades if pd.notna(g)]
            if not grades:
                return "—"
            # Prioritetas: raudona > žalia > geltona
            if any(g < 4 for g in grades):
                return "Raudona"
            if all(g >= 9 for g in grades):
                return "Žalia"
            if min(grades) >= 7:
                return "Geltona"
            return "—"


        summary_df = grade_lists.copy()
        summary_df["Žyma"] = summary_df["Visi_iverciai"].apply(student_flag)
        summary_df["Kiek įvertinimų"] = summary_df["Visi_iverciai"].apply(len)
        summary_df["Min"] = summary_df["Visi_iverciai"].apply(min)
        summary_df["Max"] = summary_df["Visi_iverciai"].apply(max)
        summary_df["Vidurkis"] = summary_df["Visi_iverciai"].apply(
            lambda x: sum(x) / len(x) if len(x) else None
        )

        # Pervadiname ir sutvarkome stulpelius
        summary_df = summary_df.rename(columns={"Vardas_Pavarde": "Mokinys"})
        summary_df = summary_df[["Mokinys", "Kiek įvertinimų", "Min", "Max", "Vidurkis", "Žyma"]].sort_values(
            by=["Žyma", "Vidurkis"], ascending=[True, False]
        )


        # Eilučių spalvinimas
        def highlight_row(row):
            color_map = {
                "Žalia": "#c7f1c0",  # green-ish
                "Geltona": "#fff3b0",  # yellow-ish
                "Raudona": "#f8c8c8",  # red-ish
            }
            bg = color_map.get(row["Žyma"], "")
            return [f"background-color: {bg}"] * len(row)


        styled_summary = (
            summary_df.style
            .apply(highlight_row, axis=1)
            .format({"Vidurkis": "{:.2f}"})
        )

        st.dataframe(styled_summary, use_container_width=True)

        with st.expander("Legenda"):
            st.markdown(
                "- 🟩 **Žalia** – visi pažymiai **9–10**.\n"
                "- 🟨 **Geltona** – visi pažymiai **7–10** (bet ne visi 9–10).\n"
                "- 🟥 **Raudona** – yra bent vienas pažymys **<4**.\n"
                "- —  – taisyklės netaikomos (mišri situacija)."
            )

    # === 📊 LYGIAI
    with tab_levels:
        st.subheader("Pasiekimų lygių pasiskirstymas (pagal mažiausią pažymį)")

        per_student_grades = filtered_df.groupby("Vardas_Pavarde")["Ivertinimas"].apply(list)
        levels_series = per_student_grades.apply(student_overall_level)

        level_order = [
            "Aukštesnysis (9–10)",
            "Pagrindinis (7–10)",
            "Patenkinamas (5–10)",
            "Slenkstinis (4–10)",
            "Nepatenkinamas (1–10)"
        ]
        level_counts = levels_series.value_counts().reindex(level_order, fill_value=0).reset_index()
        level_counts.columns = ["Lygis", "Kiekis"]

        fig_levels = px.bar(
            level_counts, x="Lygis", y="Kiekis", text="Kiekis",
            color="Lygis", color_discrete_map=level_color_map,
            title="Mokinių skaičius pagal pasiekimų lygį"
        )
        fig_levels.update_traces(textposition="outside")
        fig_levels.update_layout(xaxis_tickangle=-10, margin=dict(l=10, r=20, t=50, b=10), height=500, showlegend=False)
        st.plotly_chart(fig_levels, use_container_width=True)

        st.dataframe(level_counts, use_container_width=True)

    # === 📤 EKSPORTAS (Excel + PDF)
    with tab_exports:
        st.subheader("Eksportas")

        # — Spalvintas Excel (vardų stulpelis)
        st.markdown("##### 📥 Atsisiųsti Excel su nuspalvintais vardais")

        # Žymų -> ARGB spalvų žemėlapis (8-ženklis ARGB: FF + RGB)
        color_map_argb = {
            "Aukštesnysis (9–10)": "FF86EFAC",  # dėstom spalvas pagal lygius, bet vardams toliau taikysime trišakį
            "Pagrindinis (7–10)":  "FFFDE68A",
            "Patenkinamas (5–10)": "FFBFDBFE",
            "Slenkstinis (4–10)":  "FFFECDD3",
            "Nepatenkinamas (1–10)": "FFFCA5A5",
            # Atgalinis suderinimas su ankstesniu trišakiu:
            "Žalia":  "FFC7F1C0",
            "Geltona":"FFFFF3B0",
            "Raudona":"FFF8C8C8",
        }

        # Santraukai spalvinimui (žalia/geltona/raudona pagal jūsų taisykles)
        grade_lists = df_long.groupby("Vardas_Pavarde")["Ivertinimas"].apply(list).reset_index(name="Visi_iverciai")
        def student_flag(grades):
            grades = [g for g in grades if pd.notna(g)]
            if not grades: return "—"
            if any(g < 4 for g in grades): return "Raudona"
            if all(g >= 9 for g in grades): return "Žalia"
            if min(grades) >= 7: return "Geltona"
            return "—"

        summary_df = grade_lists.copy()
        summary_df["Žyma"] = summary_df["Visi_iverciai"].apply(student_flag)
        summary_df = summary_df.rename(columns={"Vardas_Pavarde": "Mokinys"})

        def norm_name(s: str) -> str:
            if s is None: return ""
            s = str(s).replace("\u00A0", " ").strip()
            s = " ".join(s.split())
            return s.casefold()

        student_to_fill = {
            norm_name(row["Mokinys"]): color_map_argb.get(row["Žyma"], None)
            for _, row in summary_df.iterrows()
        }

        from openpyxl.utils import column_index_from_string
        def find_name_column(ws, header_row_idx=4):
            possible_headers = {
                "vardas", "vardas pavardė", "vardas ir pavardė",
                "vardas_pavarde", "vardas, pavardė", "vardas_pavardė"
            }
            try:
                header_row = ws[header_row_idx]
                for cell in header_row:
                    val = str(cell.value).strip() if cell.value is not None else ""
                    if norm_name(val) in possible_headers:
                        col = cell.column
                        if isinstance(col, str):
                            return column_index_from_string(col)
                        return int(col)
            except Exception:
                pass
            return 2  # B kaip numatytasis

        def colorize_names_in_workbook(xlsx_bytes: bytes, sheet_name="Pasiekimų ir lankomumo"):
            wb = load_workbook(BytesIO(xlsx_bytes))
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
            name_col = find_name_column(ws, header_row_idx=4)

            def make_fill(argb):
                return PatternFill(fill_type="solid", fgColor=argb) if argb else None

            start_row = 5
            for row_idx in range(start_row, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=name_col)
                raw = cell.value
                if raw is None:
                    continue
                key = norm_name(raw)
                argb = student_to_fill.get(key)
                if argb:
                    cell.fill = make_fill(argb)

            out = BytesIO()
            wb.save(out)
            out.seek(0)
            return out

        if st.button("📗 Generuoti Excel su nuspalvintais vardais"):
            colored_excel = colorize_names_in_workbook(uploaded_bytes, sheet_name="Pasiekimų ir lankomumo")
            st.success("✅ Excel paruoštas.")
            st.download_button(
                "📥 Atsisiųsti nuspalvintą Excel",
                data=colored_excel,
                file_name="su_spalvomis.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        st.divider()

        # — PDF ataskaita (naudojame STATIC matplotlib figūras PDF’ui)
        st.markdown("##### 📄 Generuoti PDF ataskaitą")

        FONT_PATH = "DejaVuSans.ttf"
        pdfmetrics.registerFont(TTFont("DejaVu", FONT_PATH))

        def build_matplotlib_figs_for_pdf():
            # subject_avg (iš „overview“)
            subject_avg_local = filtered_df.groupby("Dalykas")["Ivertinimas"].mean().sort_values()
            fig1, ax1 = plt.subplots()
            ax1.barh(subject_avg_local.index, subject_avg_local.values)
            ax1.set_xlabel("Vidurkis"); ax1.set_xlim(0, 10); ax1.set_title("Vidutiniai įvertinimai pagal dalyką")
            for bar in ax1.patches:
                w = bar.get_width()
                ax1.text(w + 0.1, bar.get_y() + bar.get_height()/2, f"{w:.1f}", va='center', fontsize=8)

            # student_avg
            student_avg_local = filtered_df.groupby("Vardas_Pavarde")["Ivertinimas"].mean().sort_values(ascending=True)
            fig2, ax2 = plt.subplots()
            ax2.barh(student_avg_local.index, student_avg_local.values)
            ax2.set_xlabel("Vidurkis"); ax2.set_xlim(0, 10); ax2.set_title("Mokiniai pagal bendrą vidurkį")
            for bar in ax2.patches:
                w = bar.get_width()
                ax2.text(w + 0.1, bar.get_y() + bar.get_height()/2, f"{w:.1f}", va='center', fontsize=8)

            # levels
            per_student = filtered_df.groupby("Vardas_Pavarde")["Ivertinimas"].apply(list)
            levels_pdf = per_student.apply(student_overall_level).value_counts().sort_index()
            fig3, ax3 = plt.subplots()
            ax3.bar(levels_pdf.index, levels_pdf.values)
            ax3.set_ylabel("Mokinių skaičius"); ax3.set_title("Pasiekimų lygių pasiskirstymas")
            for idx, v in enumerate(levels_pdf.values):
                ax3.text(idx, v + 0.05, str(int(v)), ha='center', va='bottom', fontsize=8)
            plt.xticks(rotation=15)
            return [fig1, fig2, fig3]

        def generate_recommendations(student_df):
            avg = student_df["Ivertinimas"].mean()
            if avg >= 9: return "Puikūs rezultatai. Rekomenduojame tęsti tokiu pačiu tempu."
            if avg >= 7: return "Geri pasiekimai. Galima skirti daugiau dėmesio sunkesniems dalykams."
            if avg >= 5: return "Vidutiniai rezultatai. Rekomenduojame papildomas konsultacijas."
            return "Silpni rezultatai. Būtina stipri pagalba ir dažnesnės konsultacijos."

        if st.button("📄 Generuoti PDF ataskaitą"):
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                doc = SimpleDocTemplate(tmp.name, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm,
                                        topMargin=2*cm, bottomMargin=2*cm)
                styles = getSampleStyleSheet()
                header_style = ParagraphStyle(name='Header', fontSize=18, leading=22, spaceAfter=12, fontName='DejaVu')
                normal_style = ParagraphStyle(name='NormalLT', fontName='DejaVu', fontSize=10, leading=12)
                bold_style = ParagraphStyle(name='BoldLT', fontName='DejaVu', fontSize=10, leading=12,
                                            spaceAfter=6, textColor=colors.black, bulletFontName='DejaVu')

                story = []
                story.append(Paragraph("📘 Mokinių pasiekimų ataskaita", header_style))

                # Jei vienas mokinys pasirinktas šoniniame filtre (ne „Visi“)
                only_one_student = (student_count == 1)
                if only_one_student:
                    student_name = filtered_df["Vardas_Pavarde"].iloc[0]
                    student_df = filtered_df[filtered_df["Vardas_Pavarde"] == student_name]
                    story.append(Paragraph(f"<b>Mokinys:</b> {student_name}", normal_style))
                    data = [["Dalykas", "Įvertinimas"]] + [[row["Dalykas"], str(row["Ivertinimas"])]
                                                           for _, row in student_df.iterrows()]
                    table = Table(data, hAlign='LEFT')
                    table.setStyle(TableStyle([
                        ('FONTNAME', (0, 0), (-1, -1), 'DejaVu'),
                        ('BACKGROUND', (0, 0), (-1, 0), colors.lightblue),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                        ('FONTNAME', (0, 0), (-1, 0), 'DejaVu')
                    ]))
                    story.append(table)
                    story.append(Spacer(1, 0.5*cm))
                else:
                    story.append(Paragraph("<b>Kiekvieno dalyko klasės vidurkiai:</b>", bold_style))
                    subject_avg_df = (filtered_df.groupby("Dalykas")["Ivertinimas"]
                                      .mean().sort_values(ascending=False).reset_index())
                    subject_avg_table = [["Dalykas", "Vidurkis"]] + [
                        [row["Dalykas"], f"{row['Ivertinimas']:.2f}"] for _, row in subject_avg_df.iterrows()
                    ]
                    table = Table(subject_avg_table, hAlign='LEFT')
                    table.setStyle(TableStyle([
                        ('FONTNAME', (0, 0), (-1, -1), 'DejaVu'),
                        ('BACKGROUND', (0, 0), (-1, 0), colors.lightblue),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                        ('FONTNAME', (0, 0), (-1, 0), 'DejaVu')
                    ]))
                    story.append(table)
                    story.append(Spacer(1, 0.5*cm))

                # Įkeliame statinius grafikus PDF‘ui
                fig_list = build_matplotlib_figs_for_pdf()
                fig_paths = []
                for i, fig in enumerate(fig_list):
                    fig_path = os.path.join(tempfile.gettempdir(), f"fig_pdf_{i}.png")
                    fig.savefig(fig_path, bbox_inches="tight", dpi=300)
                    plt.close(fig)
                    fig_paths.append(fig_path)

                for path in fig_paths:
                    img = PILImage.open(path)
                    iw, ih = img.size
                    aspect = ih / iw
                    target_width = 16 * cm
                    target_height = target_width * aspect
                    story.append(Image(path, width=target_width, height=target_height))
                    story.append(Spacer(1, 0.5 * cm))

                doc.build(story)

                st.success("✅ PDF ataskaita sugeneruota")
                st.download_button("📥 Atsisiųsti PDF", data=open(tmp.name, "rb"), file_name="ataskaita.pdf")


